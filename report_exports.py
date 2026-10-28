"""Build downloadable Word and Excel reports from a real prediction report."""
from __future__ import annotations

from io import BytesIO

import pandas as pd
from docx import Document
from docx.shared import Inches


def _summary(report) -> pd.DataFrame:
    c, p, b = report.config, report.prediction, report.bottleneck
    rows = [
        ("CPU", c.cpu), ("GPU", c.gpu), ("RAM (GB)", c.ram_gb), ("VRAM (GB)", c.vram_gb),
        ("Game", c.game), ("Resolution", c.resolution), ("Ray tracing", "On" if c.ray_tracing else "Off"),
        ("Target FPS", report.target_fps), ("Predicted FPS", round(p.fps, 2)),
        ("Prediction interval low FPS", round(p.low, 2)), ("Prediction interval high FPS", round(p.high, 2)),
        ("Interval coverage", f"{p.coverage:.0%}"), ("Playability", report.tier.name),
        ("Performance level", report.level), ("Benchmark percentile", round(report.percentile, 1)),
        ("Target confidence label", report.confidence_label), ("Target confidence", f"{report.confidence:.0%}"),
        ("Advisor goal", report.advisor_goal),
        ("Advisor summary", report.advisor_summary or "Generated when prediction is displayed"),
        ("Probability of meeting target", f"{report.p_meet_target:.0%}"),
        ("Recommended graphics preset", report.quality.level), ("Recommendation basis", "Rule-based, not an ML quality classifier"),
        ("Hardware balance", b.state), ("CPU/GPU score gap (%)", round(b.gap_pct, 2)),
    ]
    return pd.DataFrame(rows, columns=["Metric", "Value"])


def build_excel(report) -> bytes:
    """Export summary, explanations, upgrade scenarios, and recommendations."""
    buffer = BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        _summary(report).to_excel(writer, sheet_name="Prediction Summary", index=False)
        report.explanation.contributions.to_excel(writer, sheet_name="AI Explanation", index=False)
        (report.upgrades if report.upgrades is not None else pd.DataFrame()).to_excel(writer, sheet_name="Upgrade Scenarios", index=False)
        pd.DataFrame([{"Type": r.kind, "Source": r.source, "Recommendation": r.title, "Details": r.body}
                      for r in report.recommendations]).to_excel(writer, sheet_name="Recommendations", index=False)
        pd.DataFrame(report.quality.trace).to_excel(writer, sheet_name="Graphics Settings Logic", index=False)
        for ws in writer.book.worksheets:
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions
            for column_cells in ws.columns:
                values = [len(str(cell.value or "")) for cell in column_cells]
                ws.column_dimensions[column_cells[0].column_letter].width = min(max(max(values, default=10) + 2, 12), 55)
    return buffer.getvalue()


def build_word(report) -> bytes:
    """Export a clean, editable Word report based on the current prediction."""
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    doc.add_heading("AI Game Performance Predictor", 0)
    doc.add_paragraph("Prediction report — generated from the current model output and selected configuration.")
    doc.add_heading("Hardware and game configuration", level=1)
    table = doc.add_table(rows=0, cols=2)
    table.style = "Light Shading Accent 1"
    for metric, value in _summary(report).itertuples(index=False, name=None):
        cells = table.add_row().cells
        cells[0].text = str(metric)
        cells[1].text = str(value)
    doc.add_heading("Prediction explanation", level=1)
    doc.add_paragraph("Contributions show how each input shifts the prediction relative to the model baseline; they are not causal guarantees.")
    for row in report.explanation.contributions.itertuples(index=False):
        doc.add_paragraph(f"{row.Feature}: {row.Contribution:+.2f} FPS ({row.Impact})", style="List Bullet")
    doc.add_heading("Hardware upgrade scenarios", level=1)
    if report.upgrades is not None and not report.upgrades.empty:
        table = doc.add_table(rows=1, cols=4)
        table.style = "Light Shading Accent 1"
        for cell, value in zip(table.rows[0].cells, ["Component", "Upgrade to", "Predicted FPS", "Change (FPS)"]):
            cell.text = value
        for _, row in report.upgrades.iterrows():
            cells = table.add_row().cells
            cells[0].text = str(row.get("Component", ""))
            cells[1].text = str(row.get("Upgrade to", ""))
            cells[2].text = f"{float(row.get('Predicted FPS', 0)):.1f}"
            cells[3].text = f"{float(row.get('Change (FPS)', 0)):+.1f}"
    else:
        doc.add_paragraph("No higher-scoring CPU or GPU options were available in the dataset catalogue.")
    doc.add_heading("AI auto-optimization advisor", level=1)
    doc.add_paragraph(f"Goal: {report.advisor_goal}")
    doc.add_paragraph(report.advisor_summary or "Advisor summary was not generated for this report.")
    doc.add_paragraph("FPS and interval are model outputs; graphics presets and hardware balance are rule-based; upgrade deltas are model estimates, not guaranteed gains.")
    doc.add_heading("Recommendations", level=1)
    for item in report.recommendations:
        doc.add_heading(f"{item.title} ({item.source})", level=2)
        doc.add_paragraph(item.body)
    doc.add_heading("Important model limitation", level=1)
    doc.add_paragraph("The benchmark dataset shows limited signal beyond GPU score and may be synthetic. Game-specific differences, settings effects, and upgrade gains should be treated as estimates, not guarantees. The model should not be assumed to generalize to hardware or settings outside its training catalogue.")
    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
